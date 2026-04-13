from flask import Flask, render_template, request, redirect, url_for, session, send_file
from flask_cors import CORS
from werkzeug.security import generate_password_hash, check_password_hash
import sqlite3
import os
from io import BytesIO
from datetime import datetime
from werkzeug.utils import secure_filename
import importlib

try:
    psycopg2 = importlib.import_module("psycopg2")
except ImportError:
    psycopg2 = None


DB_PATH = os.getenv("DATABASE_PATH", "users.db")


def get_database_url():
    return os.getenv("DATABASE_URL")


def using_postgres():
    return bool(get_database_url())


def get_db_connection():
    database_url = get_database_url()
    if database_url:
        if psycopg2 is None:
            raise RuntimeError("psycopg2-binary is required when DATABASE_URL is set.")
        sslmode = os.getenv("PGSSLMODE")
        if sslmode:
            return psycopg2.connect(database_url, sslmode=sslmode)
        return psycopg2.connect(database_url)
    return sqlite3.connect(DB_PATH)


def db_placeholders(count):
    return ", ".join(["%s" if using_postgres() else "?"] * count)


def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()

    if using_postgres():
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id SERIAL PRIMARY KEY,
            username TEXT NOT NULL UNIQUE,
            password TEXT NOT NULL,
            age INTEGER,
            grade_level TEXT,
            is_admin INTEGER NOT NULL DEFAULT 0
        )
        """)

        cursor.execute("""
        CREATE TABLE IF NOT EXISTS volunteer_hours (
            id SERIAL PRIMARY KEY,
            username TEXT NOT NULL,
            hours REAL NOT NULL,
            task TEXT NOT NULL,
            date TEXT NOT NULL,
            image TEXT
        )
        """)

        cursor.execute(
            "SELECT column_name FROM information_schema.columns WHERE table_name = 'users'"
        )
        existing_columns = {row[0] for row in cursor.fetchall()}
        if "age" not in existing_columns:
            cursor.execute("ALTER TABLE users ADD COLUMN age INTEGER")
        if "grade_level" not in existing_columns:
            cursor.execute("ALTER TABLE users ADD COLUMN grade_level TEXT")
    else:
        cursor.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            password TEXT NOT NULL,
            age INTEGER,
            grade_level TEXT,
            is_admin INTEGER NOT NULL DEFAULT 0
        )
        """)

        cursor.execute("PRAGMA table_info(users)")
        existing_columns = {row[1] for row in cursor.fetchall()}
        if "age" not in existing_columns:
            cursor.execute("ALTER TABLE users ADD COLUMN age INTEGER")
        if "grade_level" not in existing_columns:
            cursor.execute("ALTER TABLE users ADD COLUMN grade_level TEXT")

        cursor.execute("""
        CREATE TABLE IF NOT EXISTS volunteer_hours (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            hours REAL NOT NULL,
            task TEXT NOT NULL,
            date TEXT NOT NULL,
            image TEXT
        )
        """)
    conn.commit()
    conn.close()
     
app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY") or "dev-secret-change-me"
app.url_map.strict_slashes = False

# Gunicorn imports the module and does not execute the __main__ block.
# Initialize DB schema at import time so first request does not fail.
init_db()

frontend_origins = [
    origin.strip()
    for origin in os.getenv("FRONTEND_ORIGIN", "").split(",")
    if origin.strip()
]
CORS(
    app,
    resources={r"/*": {"origins": frontend_origins or "*"}},
    supports_credentials=True,
)

if frontend_origins:
    app.config["SESSION_COOKIE_SAMESITE"] = "None"
    app.config["SESSION_COOKIE_SECURE"] = True

UPLOAD_FOLDER = "static/uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
os.makedirs(UPLOAD_FOLDER, exist_ok=True)


def has_first_and_last_name(name):
    parts = [part for part in name.split() if part]
    return len(parts) >= 2


def normalize_username(name):
    return " ".join((name or "").split()).strip().lower()


def get_admin_usernames():
    admin_usernames = {
        normalize_username(os.getenv("ADMIN_USERNAME", "")),
    }
    admin_usernames.update(
        normalize_username(username)
        for username in os.getenv("ADMIN_USERNAMES", "").split(",")
        if username.strip()
    )
    return {username for username in admin_usernames if username}


def get_reserved_usernames():
    reserved_usernames = {
        normalize_username("Dawit Yemane"),
    }
    reserved_usernames.update(
        normalize_username(username)
        for username in os.getenv("RESERVED_USERNAMES", "").split(",")
        if username.strip()
    )
    return {username for username in reserved_usernames if username}


def user_is_admin(cursor, username):
    cursor.execute(
        f"SELECT is_admin FROM users WHERE username = {db_placeholders(1)}",
        (username,)
    )
    result = cursor.fetchone()
    if result and result[0] == 1:
        return True
    return normalize_username(username) in get_admin_usernames()


@app.after_request
def add_no_cache_headers(response):
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response

@app.route("/")
def root():
    return redirect(url_for("home"))


@app.route("/home")
def home():
    return render_template("home.html")


@app.route("/health")
def health():
    return {"status": "ok"}, 200

@app.route("/signup", methods=["GET", "POST"])
def signup():
    if request.method == "POST":
        username = " ".join(request.form["username"].split())
        password = request.form["password"]
        age_raw = request.form["age"].strip()
        grade_level = request.form["grade_level"].strip()

        try:
            age = int(age_raw)
        except ValueError:
            age = None

        if not username:
            return render_template(
                "signup.html",
                error="First and last name cannot be empty.",
                signup_username=username,
                signup_age=age_raw,
                signup_grade_level=grade_level
            )

        if not has_first_and_last_name(username):
            return render_template(
                "signup.html",
                error="Please enter both first and last name.",
                signup_username=username,
                signup_age=age_raw,
                signup_grade_level=grade_level
            )

        normalized_username = normalize_username(username)
        if normalized_username in get_reserved_usernames() and normalized_username not in get_admin_usernames():
            return render_template(
                "signup.html",
                error="That username is reserved and cannot be used.",
                signup_username=username,
                signup_age=age_raw,
                signup_grade_level=grade_level
            )

        if age is None or age <= 0:
            return render_template(
                "signup.html",
                error="Please enter a valid age.",
                signup_username=username,
                signup_age=age_raw,
                signup_grade_level=grade_level
            )

        if not grade_level:
            return render_template(
                "signup.html",
                error="Please enter a grade level.",
                signup_username=username,
                signup_age=age_raw,
                signup_grade_level=grade_level
            )

        hashed_password = generate_password_hash(password)
        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT username FROM users")
        existing_user = next(
            (row for row in cursor.fetchall() if normalize_username(row[0]) == normalized_username),
            None,
        )
        if existing_user:
            conn.close()
            return render_template(
                "signup.html",
                error="That name is already registered. Please log in.",
                signup_username=username,
                signup_age=age_raw,
                signup_grade_level=grade_level
            )
        
        # Check if this is the first user
        cursor.execute("SELECT COUNT(*) FROM users")
        user_count = cursor.fetchone()[0]
        is_admin = 1 if user_count == 0 else 0
        
        cursor.execute(
            f"INSERT INTO users (username, password, age, grade_level, is_admin) VALUES ({db_placeholders(5)})",
            (username, hashed_password, age, grade_level, is_admin)
        )
        conn.commit()
        conn.close()
        return redirect(url_for("login"))
    return render_template("signup.html", signup_username="", signup_age="", signup_grade_level="")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method =="POST":
        username = " ".join(request.form["username"].split())
        password = request.form["password"]

        if not username or not password:
            return render_template(
                "login.html",
                error="Please enter both username and password.",
                login_username=username
            )
        
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT username, password FROM users")
        normalized_username = normalize_username(username)
        user = next(
            (row for row in cursor.fetchall() if normalize_username(row[0]) == normalized_username),
            None,
        )
        conn.close()
        
        if user and check_password_hash(user[1], password):
            session["username"] = user[0]
            return redirect(url_for("dashboard"))
        else:
            return render_template(
                "login.html",
                error="Invalid username or password.",
                login_username=username
            )
        
    return render_template("login.html", login_username="")

@app.route("/dashboard", methods=["GET", "POST"])
def dashboard():
    if "username" not in session:
        return redirect(url_for("login"))

    username = session["username"]

    conn = get_db_connection()
    cursor = conn.cursor()

    is_admin = user_is_admin(cursor, username)

    if request.method == "POST":
        hours = request.form["hours"]
        task = request.form["task"]
        date = request.form["date"]

        image_file = request.files.get("image")

        filename = None
        if image_file and image_file.filename != "":
            filename = secure_filename(image_file.filename)
            image_path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
            image_file.save(image_path)

        cursor.execute(
            f"INSERT INTO volunteer_hours (username, hours, task, date, image) VALUES ({db_placeholders(5)})",
            (username, hours, task, date, filename)
        )
        conn.commit()

    if is_admin:
        cursor.execute(
            "SELECT username, hours, task, date, image FROM volunteer_hours ORDER BY date DESC"
        )
    else:
        cursor.execute(
            f"SELECT username, hours, task, date, image FROM volunteer_hours WHERE username={db_placeholders(1)} ORDER BY date DESC",
            (username,)
        )
    records = cursor.fetchall()

    conn.close()

    return render_template(
        "dashboard.html",
        username=username,
        records=records,
        is_admin=is_admin,
        show_all_submissions=is_admin
    )

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.route("/admin")
def admin():
    if "username" not in session:
        return redirect(url_for("login"))

    conn = get_db_connection()
    cursor = conn.cursor()

    if not user_is_admin(cursor, session["username"]):
        conn.close()
        return "Access denied"

    cursor.execute("""
        SELECT vh.username, u.age, u.grade_level, vh.hours, vh.task, vh.date, vh.image
        FROM volunteer_hours vh
        LEFT JOIN users u ON LOWER(u.username) = LOWER(vh.username)
        ORDER BY date DESC
    """)
    all_records = cursor.fetchall()

    cursor.execute(
        "SELECT username, age, grade_level, is_admin FROM users ORDER BY username"
    )
    users = cursor.fetchall()

    conn.close()

    return render_template("admin.html", records=all_records, users=users)

@app.route("/toggle_admin/<username>")
def toggle_admin(username):
    if "username" not in session:
        return redirect(url_for("login"))

    conn = get_db_connection()
    cursor = conn.cursor()

    # check current user is admin
    if not user_is_admin(cursor, session["username"]):
        conn.close()
        return "Access denied"

    # ❗ prevent removing yourself
    if username == session["username"]:
        conn.close()
        return "You cannot change your own admin status"

    # toggle admin
    cursor.execute(
        f"UPDATE users SET is_admin = CASE WHEN is_admin = 1 THEN 0 ELSE 1 END WHERE username = {db_placeholders(1)}",
        (username,)
    )

    conn.commit()
    conn.close()

    return redirect(url_for("admin"))


@app.route("/admin/export")
def export_volunteer_data():
    if "username" not in session:
        return redirect(url_for("login"))

    conn = get_db_connection()
    cursor = conn.cursor()

    if not user_is_admin(cursor, session["username"]):
        conn.close()
        return "Access denied"

    cursor.execute(
        """
        SELECT vh.username, u.age, u.grade_level, vh.hours, vh.task, vh.date, vh.image
        FROM volunteer_hours vh
        LEFT JOIN users u ON LOWER(u.username) = LOWER(vh.username)
        ORDER BY vh.date DESC
        """
    )
    records = cursor.fetchall()
    conn.close()

    try:
        from openpyxl import Workbook
    except ImportError:
        return "Excel export requires openpyxl. Install it with: pip install openpyxl"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Volunteer Submissions"

    headers = ["Username", "Age", "Grade Level", "Hours", "Task", "Date", "Image File"]
    worksheet.append(headers)

    for row in records:
        worksheet.append(list(row))

    for column in worksheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            cell_value = "" if cell.value is None else str(cell.value)
            if len(cell_value) > max_length:
                max_length = len(cell_value)
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 40)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    file_name = f"volunteer_submissions_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=file_name,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

if __name__ == "__main__":
    app.run(debug = True)
